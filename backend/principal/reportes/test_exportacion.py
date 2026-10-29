from datetime import timedelta
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook

from chat.models import Mensaje
from sesiones.services import iniciar_programa, iniciar_conversacion
from usuarios.models import Perfil
from .models import Evidencia, ReporteSemanal
from .exportacion import crear_excel
from .indicadores import perfil_indicadores


class ExportacionTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_user(username='admin', is_staff=True)
        self.usuario = User.objects.create_user(username='estudiante', first_name='Nombre privado', email='privado@gmail.com')
        self.otro = User.objects.create_user(username='otro', first_name='Otro')
        self.perfil = Perfil.objects.create(usuario=self.usuario, origen_datos='demostracion')
        self.ahora = timezone.now()
        iniciar_programa(self.perfil, self.ahora - timedelta(days=30))
        self.semanas = {s.numero: s for s in self.usuario.semanas.all()}
        self.ruta = f'/api/admin/evaluacion.xlsx?alumnos={self.usuario.pk}'
        self.client.force_login(self.admin)

    def mensaje(self, numero, texto, estado='listo', cita=None, indicador='planificacion', senal='recurso'):
        sesion = self.semanas[numero]
        m = Mensaje.objects.create(usuario=self.usuario, sesion=sesion, contenido=texto, remitente='usuario', fecha=sesion.inicio + timedelta(hours=1), analisis_estado=estado)
        if cita:
            Evidencia.objects.create(mensaje=m, indicador=indicador, cita=cita, observacion='Relato del estudiante.', senal=senal)
        return m

    def libro(self):
        respuesta = self.client.get(self.ruta)
        self.assertEqual(respuesta.status_code, 200)
        self.assertIn('no-store', respuesta.headers['Cache-Control'])
        self.assertIn('attachment;', respuesta.headers['Content-Disposition'])
        return load_workbook(BytesIO(respuesta.content))

    def valor(self, libro, indicador):
        filas = list(libro['Resultados'].values)
        return filas[1][filas[0].index(f'Análisis · {indicador}')]

    def test_exporta_id_y_siete_analisis_con_evidencias_sin_datos_personales(self):
        self.mensaje(1, 'Esta semana casi siempre organizo mis tareas.', cita='casi siempre organizo mis tareas')
        self.mensaje(4, 'Esta semana casi siempre organizo mis tareas.', cita='casi siempre organizo mis tareas')
        with patch('reportes.services.analizar_mensaje') as modelo:
            libro = self.libro()
            modelo.assert_not_called()
        self.assertEqual(libro.sheetnames, ['Resultados', 'Evidencias', 'Ficha'])
        self.assertEqual(libro['Resultados'].max_column, 9)
        self.assertEqual(libro['Resultados']['A2'].value, f'Estudiante {self.usuario.pk}')
        self.assertEqual(libro['Resultados']['B2'].value, self.usuario.username)
        self.assertIn('Recurso utilizado', self.valor(libro, 'Planificación'))
        self.assertIn('casi siempre organizo mis tareas', self.valor(libro, 'Planificación'))
        self.assertEqual(libro['Evidencias'].max_row, 3)
        contenido = str([list(h.values) for h in libro])
        self.assertNotIn(self.usuario.first_name, contenido)
        self.assertNotIn(self.usuario.email, contenido)
        self.assertNotIn('Inicial · S1', contenido)
        self.assertNotIn('Final · S4', contenido)
        self.assertEqual(ReporteSemanal.objects.count(), 0)

    def test_procedencia_solo_en_ficha_y_solo_si_hay_registros_generados(self):
        libro = self.libro()
        self.assertNotIn('Demostración', str(list(libro['Resultados'].values)))
        self.assertIn('Demostración', str(list(libro['Ficha'].values)))
        self.perfil.origen_datos = 'conversacion'
        self.perfil.save(update_fields=['origen_datos'])
        libro = self.libro()
        contenido = str([list(h.values) for h in libro]).casefold()
        self.assertNotIn('demostración', contenido)
        self.assertNotIn('fictici', contenido)

    def test_periodo_abierto_exporta_evidencia_sin_exigir_semanas_o_cierre(self):
        sesion = iniciar_conversacion(self.perfil, self.ahora - timedelta(hours=1))
        self.assertIsNone(sesion.fin)
        texto = 'Esta semana casi siempre organizo mis tareas.'
        mensaje = Mensaje.objects.create(
            usuario=self.usuario, sesion=sesion, contenido=texto,
            remitente='usuario', fecha=self.ahora - timedelta(minutes=5), analisis_estado='listo',
        )
        Evidencia.objects.create(mensaje=mensaje, indicador='planificacion', cita=texto,
                                observacion='Frecuencia explícita del estudiante.', senal='recurso')
        libro = self.libro()
        self.assertEqual(libro['Resultados'].max_column, 9)
        self.assertIn('Recurso utilizado', self.valor(libro, 'Planificación'))
        self.assertIn(texto, str(list(libro['Evidencias'].values)))

    def test_pendientes_conservan_la_evidencia_validada_y_advierten_descripcion_provisional(self):
        self.mensaje(3, 'Esta semana siempre organizo mis tareas.', cita='siempre organizo mis tareas')
        self.mensaje(4, 'Esta semana siempre organizo mis tareas.', estado='pendiente', cita='siempre organizo mis tareas')
        libro = self.libro()
        self.assertIn('Recurso utilizado', self.valor(libro, 'Planificación'))
        self.assertIn('provisional', self.valor(libro, 'Planificación'))
        self.assertEqual(libro['Evidencias'].max_row, 2)
        self.semanas[4].delete()
        self.assertNotIn('provisional', self.valor(self.libro(), 'Planificación'))

    def test_descarta_citas_inventadas_y_fechas_fuera_de_sesion(self):
        self.mensaje(1, 'Solo este texto', cita='Esta semana siempre organizo mis tareas.')
        m = self.mensaje(4, 'Esta semana siempre organizo mis tareas.', cita='siempre organizo mis tareas')
        Mensaje.objects.filter(pk=m.pk).update(fecha=self.semanas[4].fin)
        libro = self.libro()
        self.assertIn('No mencionado en la conversación', self.valor(libro, 'Planificación'))
        self.assertEqual(libro['Evidencias'].max_row, 1)

    def test_frecuencias_distintas_no_se_convierten_en_una_escala_cualitativa(self):
        for frecuencia in ('rara vez', 'siempre'):
            texto = f'Esta semana {frecuencia} organizo mis tareas.'
            self.mensaje(1, texto, cita=texto)
        libro = self.libro()
        self.assertIn('Recurso utilizado', self.valor(libro, 'Planificación'))
        self.assertEqual(libro['Evidencias'].max_row, 3)
        self.assertEqual(perfil_indicadores(self.usuario)['indicadores'][4]['frecuencia'], None)

    def test_no_exporta_formulas_desde_mensajes_y_fecha_es_lima(self):
        texto = '=HYPERLINK("https://ejemplo.invalid", "clic")'
        self.mensaje(1, texto, cita=texto)
        libro = self.libro()
        for hoja in libro:
            for fila in hoja:
                self.assertTrue(all(c.data_type != 'f' for c in fila))
        fecha = dict(libro['Ficha'].values)['Exportado']
        self.assertEqual(fecha.hour, timezone.localtime().hour)
        self.assertIsNone(fecha.tzinfo)

    def test_ningun_otro_estudiante_se_incluye(self):
        libro = self.libro()
        self.assertEqual(libro['Resultados'].max_row, 2)
        self.assertNotIn(f'Estudiante {self.otro.pk}', str(list(libro['Resultados'].values)))

    def test_anonimo_estudiante_y_seleccion_invalida_rechazados(self):
        self.client.logout()
        self.assertEqual(self.client.get(self.ruta).status_code, 401)
        self.client.force_login(self.usuario)
        self.assertEqual(self.client.get(self.ruta).status_code, 403)
        self.client.force_login(self.admin)
        for valor in ('', 'abc', '1,,2', ',1', ','.join(['1'] * 111)):
            self.assertEqual(self.client.get('/api/admin/evaluacion.xlsx', {'alumnos': valor}).status_code, 400)
        self.assertEqual(self.client.get('/api/admin/evaluacion.xlsx', {'alumnos': self.admin.pk}).status_code, 404)
        self.assertEqual(self.client.get('/api/admin/evaluacion.xlsx', {'alumnos': 9999999}).status_code, 404)

    def test_actividad_en_curso_no_inventa_una_medicion_final(self):
        m = self.mensaje(4, 'Esta semana siempre organizo mis tareas.', cita='siempre organizo mis tareas')
        ahora = m.fecha + timedelta(hours=1)
        libro = load_workbook(BytesIO(crear_excel([self.usuario], ahora=ahora)))
        self.assertIn('Recurso utilizado', self.valor(libro, 'Planificación'))
        self.assertFalse(any('Inicial' in str(celda.value) or 'Final ·' in str(celda.value) for celda in libro['Resultados'][1]))

    def test_respuesta_corta_no_crea_citas_ni_transformaciones_numericas_en_excel(self):
        padre = self.mensaje(1, 'Me falta tiempo para las tareas.')
        pregunta = Mensaje.objects.create(usuario=self.usuario, sesion=self.semanas[1], remitente='asistente', contenido='¿Con qué frecuencia te falta tiempo para las tareas esta semana?', fecha=padre.fecha + timedelta(seconds=20), respuesta_a=padre, analisis_estado='listo')
        respuesta = self.mensaje(1, 'Casi siempre.')
        Mensaje.objects.filter(pk=respuesta.pk).update(fecha=pregunta.fecha + timedelta(seconds=20))
        self.assertIn('No mencionado en la conversación', self.valor(self.libro(), 'Presión del tiempo'))
        Mensaje.objects.filter(pk=respuesta.pk).update(fecha=pregunta.fecha + timedelta(hours=3))
        self.assertIn('No mencionado en la conversación', self.valor(self.libro(), 'Presión del tiempo'))

    def test_silencio_no_se_convierte_en_ausencia_y_ninguna_celda_queda_vacia(self):
        self.mensaje(1, 'Hola, hoy tuve clases.')
        libro = self.libro()
        for celda in list(libro['Resultados'].values)[1][2:]:
            self.assertIsInstance(celda, str)
            self.assertIn('No mencionado en la conversación', celda)
            self.assertNotIn('No presenta', celda)
            self.assertNotIn('Sin información', celda)

    def test_ausencia_explicita_y_recurso_ausente_se_describen_separadamente(self):
        self.mensaje(1, 'No tengo dolores por las tareas.', cita='No tengo dolores por las tareas.',
                     indicador='sintomas_fisicos', senal='ausencia_explicita')
        self.mensaje(2, 'No pido ayuda a nadie.', cita='No pido ayuda a nadie.',
                     indicador='busqueda_soporte', senal='ausencia_explicita')
        libro = self.libro()
        self.assertIn('Ausencia expresada.', self.valor(libro, 'Síntomas físicos'))
        self.assertIn('Ausencia expresada del recurso', self.valor(libro, 'Búsqueda de soporte'))
        self.assertIn('No pido ayuda a nadie.', str(list(libro['Evidencias'].values)))

    def test_relato_mixto_conserva_citas_de_ambas_senales(self):
        self.mensaje(1, 'Me duele la cabeza por los exámenes.', cita='Me duele la cabeza por los exámenes.',
                     indicador='sintomas_fisicos', senal='dificultad')
        self.mensaje(2, 'Hoy no me duele la cabeza por estudiar.', cita='Hoy no me duele la cabeza por estudiar.',
                     indicador='sintomas_fisicos', senal='ausencia_explicita')
        libro = self.libro()
        self.assertIn('Relato mixto', self.valor(libro, 'Síntomas físicos'))
        indicador = perfil_indicadores(self.usuario)['indicadores'][2]
        self.assertEqual(indicador['estado'], 'relato_mixto')
        self.assertEqual(indicador['total_evidencias'], 2)
        self.assertEqual(len(indicador['evidencias']), 2)
        self.assertEqual(libro['Evidencias'].max_row, 3)

    def test_sin_mensajes_y_analisis_pendiente_tienen_descripcion_explicita(self):
        self.assertIn('Sin conversación registrada', self.valor(self.libro(), 'Planificación'))
        self.mensaje(1, 'Organizo mis tareas.', estado='pendiente', cita='Organizo mis tareas.')
        libro = self.libro()
        self.assertIn('Análisis pendiente', self.valor(libro, 'Planificación'))
        self.assertEqual(libro['Evidencias'].max_row, 1)
