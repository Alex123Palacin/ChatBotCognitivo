from datetime import datetime, timedelta
from io import StringIO, BytesIO
import json
from zoneinfo import ZoneInfo
import importlib
from types import SimpleNamespace

from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase, override_settings
from django.apps import apps
from django.db import connection
from openpyxl import load_workbook

from chat.models import Mensaje
from sesiones.models import SesionSemanal
from reportes.models import MuestraEstudio, Evidencia
from reportes.estudio import muestra_actual
from reportes.indicadores import perfil_indicadores


class EstudioTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.fecha = datetime(2026, 9, 11, 13, tzinfo=ZoneInfo('America/Lima'))
        User = get_user_model()
        User.objects.bulk_create([User(username=f'estudiante{i}', password='!') for i in range(136)])
        for u in User.objects.all():
            s = SesionSemanal.objects.create(usuario=u, numero=1, inicio=cls.fecha, fin=None)
            Mensaje.objects.create(usuario=u, sesion=s, contenido='Hola', remitente='usuario', fecha=cls.fecha, analisis_estado='listo')
        cls.admin = User.objects.create_user('admin', is_staff=True)

    def setUp(self):
        self.client.force_login(self.admin)

    def seleccionar(self, aplicar=True):
        salida = StringIO()
        call_command('seleccionar_estudio', aplicar=aplicar, stdout=salida)
        return json.loads(salida.getvalue())

    def test_simulacion_no_selecciona_y_muestra_se_conserva_sin_repetir_sorteo(self):
        self.assertEqual(self.seleccionar(False)['poblacion'], 136)
        self.assertFalse(MuestraEstudio.objects.exists())
        self.assertEqual(self.seleccionar()['seleccionados'], 110)
        ids = set(muestra_actual().usuarios.values_list('pk', flat=True))
        self.assertEqual(len(ids), 110)
        self.assertTrue(self.seleccionar()['existente'])
        self.assertEqual(ids, set(muestra_actual().usuarios.values_list('pk', flat=True)))
        self.assertEqual(get_user_model().objects.filter(is_staff=False).count(), 136)
        self.assertEqual(Mensaje.objects.count(), 136)

    def test_listado_excel_110_y_exclusion_sin_borrar_otras_cuentas(self):
        self.seleccionar()
        r = self.client.get('/api/admin/alumnos/').json()
        self.assertEqual(len(r['alumnos']), 110)
        self.assertEqual(r['estudio']['poblacion'], 136)
        ids = [u['id'] for u in r['alumnos']]
        respuesta = self.client.get('/api/admin/evaluacion.xlsx', {'alumnos': ','.join(map(str, ids))})
        self.assertEqual(respuesta.status_code, 200)
        w = load_workbook(BytesIO(respuesta.content))
        self.assertEqual((w['Resultados'].max_row, w['Resultados'].max_column), (111, 9))
        self.assertEqual({row[0] for row in list(w['Resultados'].values)[1:]}, {f'Estudiante {i}' for i in ids})
        fuera = get_user_model().objects.filter(is_staff=False).exclude(pk__in=ids).first()
        self.assertEqual(self.client.get('/api/admin/evaluacion.xlsx', {'alumnos': fuera.pk}).status_code, 400)
        self.assertTrue(Mensaje.objects.filter(usuario=fuera).exists())

    def test_grafico_y_excel_comparten_frecuencia_y_no_usar_recuentos_como_gravedad(self):
        self.seleccionar()
        u = muestra_actual().usuarios.first()
        m = u.mensajes_chat.get()
        m.contenido = 'Esta semana casi siempre organizo mis tareas.'; m.save()
        Evidencia.objects.create(mensaje=m, indicador='planificacion', cita=m.contenido, senal='recurso', observacion='Organización expresada.')
        fuera = Mensaje.objects.create(usuario=u, sesion=m.sesion, remitente='usuario', fecha=self.fecha + timedelta(days=1), contenido='Esta semana nunca organizo mis tareas.', analisis_estado='listo')
        Evidencia.objects.create(mensaje=fuera, indicador='planificacion', cita=fuera.contenido, senal='ausencia_explicita', observacion='Fuera de ventana.')
        p = perfil_indicadores(u)
        i = next(x for x in p['indicadores'] if x['id'] == 'planificacion')
        self.assertEqual(i['frecuencia'], 4)
        self.assertTrue(i['recurso'])
        self.assertIn('planificación', p['resumen'])
        self.assertTrue(all(x['frecuencia'] is None for x in p['indicadores'] if x['id'] != 'planificacion'))
        w = load_workbook(BytesIO(self.client.get('/api/admin/evaluacion.xlsx', {'alumnos': u.pk}).content))
        filas = list(w['Resultados'].values)
        self.assertEqual(filas[1][filas[0].index('Actividad · Planificación')], 4)

    def test_poblacion_distinta_y_muestra_incompleta_no_se_fuerzan(self):
        u = get_user_model().objects.filter(is_staff=False).first(); u.delete()
        with self.assertRaises(CommandError):
            self.seleccionar()
        self.assertFalse(MuestraEstudio.objects.exists())

    def test_sintesis_no_usa_citas_inexistentes_ni_analisis_pendiente(self):
        u = get_user_model().objects.filter(is_staff=False).first()
        m = u.mensajes_chat.get()
        Evidencia.objects.create(mensaje=m, indicador='presion_tiempo', cita='Me falta tiempo.',
                                senal='dificultad', observacion='No está en el mensaje.')
        pendiente = Mensaje.objects.create(usuario=u, sesion=m.sesion, remitente='usuario',
                                          fecha=self.fecha, contenido='Me duele la cabeza.', analisis_estado='pendiente')
        Evidencia.objects.create(mensaje=pendiente, indicador='sintomas_fisicos', cita=pendiente.contenido,
                                senal='dificultad', observacion='Aún sin analizar.')
        p = perfil_indicadores(u)
        self.assertFalse(any(i['con_evidencia'] for i in p['indicadores']))
        self.assertNotIn('se expresan dificultades', p['resumen'])
        self.assertEqual(p['analisis_pendientes'], 1)

    def test_migracion_inicializa_una_vez_sin_cambiar_registros(self):
        inicializar = importlib.import_module('reportes.migrations.0003_seleccion_inicial_estudio').seleccionar_si_procede
        anteriores = list(Mensaje.objects.values())
        inicializar(apps, SimpleNamespace(connection=connection))
        ids = set(muestra_actual().usuarios.values_list('pk', flat=True))
        self.assertEqual(len(ids), 110)
        inicializar(apps, SimpleNamespace(connection=connection))
        self.assertEqual(ids, set(muestra_actual().usuarios.values_list('pk', flat=True)))
        self.assertEqual(anteriores, list(Mensaje.objects.values()))

    @override_settings(ESTUDIO_MUESTRA_OBLIGATORIA=True)
    def test_arranque_reutiliza_los_mismos_110_e_ignora_cuentas_nuevas(self):
        antes = list(Mensaje.objects.values())
        call_command('seleccionar_estudio', aplicar=True, si_procede=True, stdout=StringIO())
        ids = set(muestra_actual().usuarios.values_list('pk', flat=True))
        get_user_model().objects.create_user('otra_cuenta_fuera_del_estudio')
        call_command('seleccionar_estudio', aplicar=True, si_procede=True, stdout=StringIO())
        listado = self.client.get('/api/admin/alumnos/')
        self.assertEqual(listado.status_code, 200)
        self.assertEqual(ids, {a['id'] for a in listado.json()['alumnos']})
        excel = self.client.get('/api/admin/evaluacion.xlsx', {'alumnos': ','.join(map(str, ids))})
        self.assertEqual(excel.status_code, 200)
        filas = list(load_workbook(BytesIO(excel.content))['Resultados'].values)[1:]
        self.assertEqual(len(filas), 110)
        self.assertEqual({f[0] for f in filas}, {f'Estudiante {pk}' for pk in ids})
        self.assertEqual(get_user_model().objects.filter(is_staff=False).count(), 137)
        self.assertEqual(antes, list(Mensaje.objects.values()))

    @override_settings(ESTUDIO_MUESTRA_OBLIGATORIA=True)
    def test_produccion_sin_muestra_no_lista_ni_exporta_otras_cuentas(self):
        self.assertEqual(self.client.get('/api/admin/alumnos/').status_code, 503)
        pk = get_user_model().objects.filter(is_staff=False).first().pk
        self.assertEqual(self.client.get('/api/admin/evaluacion.xlsx', {'alumnos': pk}).status_code, 503)
        self.assertFalse(MuestraEstudio.objects.exists())

    @override_settings(ESTUDIO_MUESTRA_OBLIGATORIA=True)
    def test_muestra_incompleta_se_detecta_y_no_se_vuelve_a_sortear(self):
        self.seleccionar()
        muestra = muestra_actual()
        muestra.usuarios.remove(muestra.usuarios.first())
        ids = set(muestra.usuarios.values_list('pk', flat=True))
        with self.assertRaises(CommandError):
            call_command('seleccionar_estudio', aplicar=True, si_procede=True, stdout=StringIO())
        self.assertEqual(self.client.get('/api/admin/alumnos/').status_code, 503)
        self.assertEqual(self.client.get('/api/admin/evaluacion.xlsx', {'alumnos': next(iter(ids))}).status_code, 503)
        self.assertEqual(ids, set(muestra.usuarios.values_list('pk', flat=True)))
        self.assertEqual(get_user_model().objects.filter(is_staff=False).count(), 136)

    def test_arranque_sin_poblacion_completa_espera_sin_completar_con_otros(self):
        get_user_model().objects.filter(is_staff=False).first().delete()
        salida = StringIO()
        call_command('seleccionar_estudio', aplicar=True, si_procede=True, stdout=salida)
        self.assertEqual(json.loads(salida.getvalue())['estado'], 'pendiente_de_poblacion')
        self.assertFalse(MuestraEstudio.objects.exists())
