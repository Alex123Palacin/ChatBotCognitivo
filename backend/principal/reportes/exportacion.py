"""Exportación del análisis cualitativo sustentado en citas de la conversación."""
from datetime import datetime
from io import BytesIO
from math import ceil
import re
from zoneinfo import ZoneInfo

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ml.analisis import INDICADORES, DIMENSIONES
from .indicadores import datos_indicadores, perfil_desde_datos
from .estudio import muestra_actual

LIMA = ZoneInfo('America/Lima')
CONTROL_XML = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def hora_lima(fecha):
    # Excel no conserva zonas horarias: la ficha identifica la zona de todas las fechas.
    return fecha.astimezone(LIMA).replace(tzinfo=None) if fecha else None


def crear_hoja(libro, nombre, columnas):
    hoja = libro.create_sheet(nombre)
    hoja.append([c[0] for c in columnas])
    hoja.sheet_view.showGridLines = False
    hoja.freeze_panes = 'C2' if len(columnas) > 2 else 'A2'
    for indice, (_, ancho) in enumerate(columnas, 1):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho
    for celda in hoja[1]:
        celda.fill = PatternFill('solid', fgColor='183653')
        celda.font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
        celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    hoja.row_dimensions[1].height = 34
    return hoja


def agregar(hoja, valores):
    hoja.append(valores)
    for celda in hoja[hoja.max_row]:
        if isinstance(celda.value, str):
            # Las conversaciones son texto, incluso si empiezan con =, +, - o @.
            celda.value = CONTROL_XML.sub('', celda.value)[:32767]
            celda.data_type = 's'
        if isinstance(celda.value, datetime):
            celda.number_format = 'dd/mm/yyyy hh:mm'
        celda.font = Font(name='Calibri', size=11, color='24374A')
        celda.alignment = Alignment(vertical='top', wrap_text=True)
        if hoja.max_row % 2 == 0:
            celda.fill = PatternFill('solid', fgColor='F0F5FA')
    lineas = max((sum(max(1, ceil(len(linea) / max(8, hoja.column_dimensions[c.column_letter].width - 3))) for linea in str(c.value or '').split('\n')) for c in hoja[hoja.max_row]), default=1)
    hoja.row_dimensions[hoja.max_row].height = min(409, max(30, lineas * 15 + 10))


def abreviar(texto, limite):
    return texto if len(texto) <= limite else texto[:limite].rstrip() + '…'


def celda_cualitativa(indicador):
    texto = indicador['etiqueta'] + '. ' + abreviar(indicador['analisis'], 900)
    if indicador['evidencias']:
        texto += '\nCita de respaldo: «' + abreviar(indicador['evidencias'][0]['cita'], 200) + '»'
    return texto


def crear_excel(usuarios, ahora=None):
    ahora = ahora or timezone.now()
    usuarios = sorted(usuarios, key=lambda usuario: usuario.pk)
    datos = datos_indicadores(usuarios, ahora)
    muestra = muestra_actual()
    periodo = (f'Mensajes del {hora_lima(muestra.inicio)} al {hora_lima(muestra.fin)} (fin exclusivo), hora de Lima. '
               'Sin dividir en semanas ni suponer una comparación inicial/final.' if muestra else
               'Actividad disponible hasta la fecha de exportación, sin dividir en semanas ni suponer una comparación inicial/final.')

    libro = Workbook()
    libro.remove(libro.active)
    libro.properties.creator = 'Chatbot Cognitivo'
    libro.properties.title = 'Análisis cualitativo de la conversación por indicador'
    columnas = [('ID estudiante', 20), ('Usuario', 25)]
    columnas.extend((f'Análisis · {meta["nombre"]}', 60) for meta in INDICADORES.values())
    resultados = crear_hoja(libro, 'Resultados', columnas)
    resultados.freeze_panes = 'C2'
    resultados.row_dimensions[1].height = 60
    evidencias = crear_hoja(libro, 'Evidencias', [
        ('ID estudiante', 20), ('Usuario', 25), ('Dimensión', 28), ('Indicador', 28),
        ('Estado cualitativo', 30), ('Cita literal del estudiante', 75),
        ('Observación registrada', 75), ('Fecha · Lima', 22),
    ])
    evidencias.freeze_panes = 'C2'
    ficha = crear_hoja(libro, 'Ficha', [('Campo', 31), ('Descripción', 115)])
    ficha.freeze_panes = 'A2'
    for fila in [
        ('Documento', 'Análisis cualitativo de la conversación por indicador · versión 4'),
        ('Exportado', hora_lima(ahora)),
        ('Periodo', periodo),
        ('Fuente', 'Mensajes del estudiante analizados y citas copiadas literalmente de esos mensajes, dentro del periodo y de su propia conversación. Las observaciones ya registradas describen esas citas; la descarga no solicita nuevas inferencias al modelo.'),
        ('Dificultad reportada', 'El estudiante expresa una dificultad relacionada con el indicador. Es una descripción del relato, no una frecuencia ni un nivel clínico.'),
        ('Recurso utilizado', 'El estudiante relata una práctica de planificación, automotivación o búsqueda de soporte que utiliza.'),
        ('Ausencia expresada', 'El estudiante niega expresamente una dificultad. En afrontamiento, significa que declara no usar o no disponer del recurso. Se limita al contexto relatado.'),
        ('Relato mixto', 'Existen distintas señales del mismo indicador. Se conservan las dificultades, recursos o ausencias expresadas sin descartar los relatos diferentes.'),
        ('No mencionado', 'El indicador no aparece en los mensajes analizados. No permite afirmar que el estudiante no presente ese problema.'),
        ('Análisis pendiente', 'Hay mensajes aún sin analizar y el indicador todavía no tiene citas validadas. Una descripción con evidencia disponible también puede ser provisional.'),
        ('Sin conversación registrada', 'No hay mensajes del estudiante dentro del periodo consultado. Cada indicador conserva una descripción explícita.'),
        ('Lectura del archivo', 'Resultados: una fila por estudiante, con ID, usuario y siete análisis cualitativos. Evidencias: citas completas, observaciones y fechas que respaldan esos análisis. Las citas abreviadas en Resultados se conservan completas en Evidencias.'),
        ('Alcance', 'Análisis cualitativo descriptivo, sin escala numérica, suma global ni clasificación de estrés bajo, medio o alto. No constituye un diagnóstico, una puntuación SISCO o una prueba de eficacia.'),
        ('Identificadores', 'ID interno y nombre de usuario. No se incluyen nombres personales, correos ni contraseñas. La hoja Evidencias contiene fragmentos de conversación y requiere acceso autorizado.'),
    ]:
        agregar(ficha, fila)
    if muestra:
        agregar(ficha, ('Selección', f'Muestra aleatoria fija de {muestra.tamano} estudiantes entre {len(muestra.poblacion_ids)}; fecha de selección: {hora_lima(muestra.creada)}. La descarga contiene los estudiantes solicitados dentro de esa muestra.'))
        agregar(ficha, ('Ventana de estudio', f'{hora_lima(muestra.inicio)} a {hora_lima(muestra.fin)} · America/Lima.'))
    demos = []
    for usuario in usuarios:
        codigo = f'Estudiante {usuario.pk}'
        perfil = getattr(usuario, 'perfil', None)
        if perfil and perfil.origen_datos == 'demostracion':
            demos.append(codigo)
        perfil_cualitativo = perfil_desde_datos(datos[usuario.pk])
        por_codigo = {i['id']: i for i in perfil_cualitativo['indicadores']}
        fila = [codigo, usuario.username]
        fila.extend(celda_cualitativa(por_codigo[indicador]) for indicador in INDICADORES)
        agregar(resultados, fila)
        for evidencia in datos[usuario.pk]['evidencias']:
            meta = INDICADORES[evidencia.indicador]
            etiqueta = {'dificultad': 'Dificultad reportada', 'recurso': 'Recurso utilizado',
                        'ausencia_explicita': 'Ausencia expresada del recurso' if meta['dimension'] == 'afrontamiento' else 'Ausencia expresada'}[evidencia.senal]
            agregar(evidencias, [codigo, usuario.username, DIMENSIONES[meta['dimension']], meta['nombre'],
                                etiqueta, evidencia.cita, evidencia.observacion.strip() or 'Cita literal del estudiante.',
                                hora_lima(evidencia.mensaje.fecha)])
    if demos:
        agregar(ficha, ('Tipo de datos', 'Demostración: ' + ', '.join(demos) + '. Sus relatos fueron generados para comprobar el sistema; no son observaciones de participantes de investigación.'))
    for hoja in libro:
        hoja.auto_filter.ref = hoja.dimensions
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()
