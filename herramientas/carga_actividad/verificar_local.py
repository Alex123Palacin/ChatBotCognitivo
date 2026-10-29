"""Prueba el paquete íntegro en SQLite aislada, sin abrir la base habitual."""
from datetime import datetime
import hashlib
from io import BytesIO, StringIO
import json
import os
from pathlib import Path
import secrets
import sys

RAIZ = Path(__file__).resolve().parents[2]


def ejecutar(salida):
    salida = Path(salida).resolve()
    salida.mkdir(parents=True, exist_ok=True)
    base_original = RAIZ / 'backend' / 'principal' / 'db.sqlite3'
    huella_original = hashlib.sha256(base_original.read_bytes()).hexdigest() if base_original.exists() else None
    base_local = salida / 'actividad_local.sqlite3'
    assert base_local.parent == salida and base_local != base_original.resolve()
    os.environ.update({'DJANGO_SETTINGS_MODULE': 'principal.settings', 'DJANGO_DB_PATH': str(base_local),
                       'DJANGO_DEBUG': 'true', 'DJANGO_ALLOWED_HOSTS': 'testserver,localhost,127.0.0.1',
                       'DJANGO_SECURE_SSL_REDIRECT': 'false', 'DJANGO_COOKIE_SECURE': 'false'})
    sys.path.insert(0, str(RAIZ / 'backend' / 'principal'))
    import django
    django.setup()
    from django.conf import settings
    from django.contrib.auth import authenticate, get_user_model
    from django.core.management import call_command
    from django.db import connection
    from django.test import Client
    from openpyxl import load_workbook
    from chat.models import Mensaje
    from reportes.models import Evidencia, ReporteSemanal
    from sesiones.models import SesionSemanal
    from usuarios.models import Perfil
    from usuarios.retencion import limpiar_conversaciones
    from reportes.services import procesar_pendientes

    assert Path(settings.DATABASES['default']['NAME']).resolve() == base_local
    call_command('migrate', interactive=False, stdout=StringIO())
    ruta = salida / 'actividad.json'
    sha = hashlib.sha256(ruta.read_bytes()).hexdigest()
    paquete = json.loads(ruta.read_text(encoding='utf-8'))
    accesos = json.loads((salida / 'accesos_privados.json').read_text(encoding='utf-8'))

    def conteos():
        return {'usuarios': get_user_model().objects.count(), 'perfiles': Perfil.objects.count(),
                'sesiones': SesionSemanal.objects.count(), 'mensajes': Mensaje.objects.count(),
                'evidencias': Evidencia.objects.count(), 'reportes': ReporteSemanal.objects.count()}

    antes = conteos()
    call_command('cargar_actividad', str(ruta), sha256=sha, stdout=StringIO())
    assert conteos() == antes, 'El modo validación escribió datos'
    print('Validación sin escrituras: correcta', flush=True)
    call_command('cargar_actividad', str(ruta), sha256=sha, aplicar=True, stdout=StringIO())
    esperado = {'mensajes': sum(len(c['mensajes']) for c in paquete['cuentas']),
                'evidencias': sum(len(m['evidencias']) for c in paquete['cuentas'] for m in c['mensajes'])}
    assert Mensaje.objects.count() == esperado['mensajes'] and Evidencia.objects.count() == esperado['evidencias']
    assert SesionSemanal.objects.count() == ReporteSemanal.objects.count() == 136
    assert not Perfil.objects.filter(terminos_aceptados_en__isnull=False).exists()
    assert not Mensaje.objects.filter(eliminar_despues_de__isnull=False).exists()
    assert not Mensaje.objects.exclude(analisis_estado='listo').exists()
    assert Perfil.objects.filter(origen_datos='demostracion').count() == 136
    despues = conteos()
    call_command('cargar_actividad', str(ruta), sha256=sha, aplicar=True, stdout=StringIO())
    assert conteos() == despues, 'Repetir la importación duplicó datos'
    print('Importación completa e idempotencia: correctas', flush=True)

    # Verificar una contraseña de cada perfil sin imprimirlas ni activar últimas actividades.
    for indice in (0, 31, 65, 99, 135):
        cuenta = accesos[indice]
        assert authenticate(username=cuenta['email'], password=cuenta['password']) is not None
    cliente = Client(enforce_csrf_checks=True)
    token = cliente.get('/api/csrf/').json()['csrfToken']
    respuesta = cliente.post('/api/auth/login/', data=json.dumps({'email': accesos[0]['email'], 'password': accesos[0]['password']}),
                             content_type='application/json', HTTP_X_CSRFTOKEN=token)
    assert respuesta.status_code == 200 and not respuesta.json()['usuario']['es_admin']
    chat = cliente.get('/api/chat/')
    assert chat.status_code == 200, chat.status_code
    assert cliente.get('/api/admin/alumnos/').status_code == 403
    print('Acceso de estudiantes, historial y permisos: correctos', flush=True)

    # Cuenta local de revisión independiente: nunca forma parte del paquete del servidor.
    admin_ruta = salida / 'acceso_admin_local.json'
    if admin_ruta.exists():
        admin_acceso = json.loads(admin_ruta.read_text(encoding='utf-8'))
        admin = get_user_model().objects.get(username=admin_acceso['usuario'])
    else:
        admin_acceso = {'usuario': 'revision.local', 'password': secrets.token_urlsafe(18) + '9!'}
        admin = get_user_model().objects.create_user(username=admin_acceso['usuario'], password=admin_acceso['password'],
                                                      first_name='Administrador', is_staff=True, is_superuser=False)
        Perfil.objects.create(usuario=admin)
        admin_ruta.write_text(json.dumps(admin_acceso, indent=2) + '\n', encoding='utf-8')
        admin_ruta.chmod(0o600)
    cliente_admin = Client()
    cliente_admin.force_login(admin)
    listado = cliente_admin.get('/api/admin/alumnos/')
    assert listado.status_code == 200 and len(listado.json()['alumnos']) == 136
    ids = [a['id'] for a in listado.json()['alumnos']]
    for uid in ids:
        detalle = cliente_admin.get(f'/api/admin/alumnos/{uid}/')
        assert detalle.status_code == 200
        reporte = detalle.json()['reporte_actual']
        assert reporte is not None and reporte['cobertura']['analisis_pendientes'] == 0
        assert len(detalle.json()['sesiones']) == 1 and detalle.json()['sesiones'][0]['fin'] is None
    print('Panel y análisis de las 136 cuentas: correctos', flush=True)
    for n, seleccion in enumerate((ids[:100], ids[100:]), 1):
        respuesta = cliente_admin.get('/api/admin/evaluacion.xlsx', {'alumnos': ','.join(map(str, seleccion))})
        assert respuesta.status_code == 200
        libro = load_workbook(BytesIO(respuesta.content), read_only=True, data_only=True)
        assert libro['Resultados'].max_row == len(seleccion) + 1 and libro['Resultados'].max_column == 8
        (salida / f'evaluacion_{n}.xlsx').write_bytes(respuesta.content)
    print('Excel en lotes de 100 y 36: correcto', flush=True)

    antes_worker = conteos()
    limpieza = limpiar_conversaciones(aplicar=True)
    procesados = procesar_pendientes(limite=5)
    assert procesados == 0 and conteos() == antes_worker
    # Login/me actualiza actividad; la segunda importación debe conservarla.
    actividad = Perfil.objects.get(usuario__email=accesos[0]['email']).ultima_actividad
    call_command('cargar_actividad', str(ruta), sha256=sha, aplicar=True, stdout=StringIO())
    assert Perfil.objects.get(usuario__email=accesos[0]['email']).ultima_actividad == actividad
    assert conteos() == antes_worker
    connection.close()
    huella_despues = hashlib.sha256(base_original.read_bytes()).hexdigest() if base_original.exists() else None
    assert huella_original == huella_despues, 'La base habitual cambió durante la prueba'
    resultado = {'base': str(base_local), 'cuentas_cargadas': 136, **esperado,
                 'sesiones': 136, 'reportes': 136, 'pruebas_acceso': 5,
                 'detalles_admin_comprobados': 136, 'exportaciones': [100, 36],
                 'idempotencia': True, 'worker_sin_pendientes': True,
                 'base_habitual_sin_cambios': True, 'fecha_verificacion': datetime.now().astimezone().isoformat()}
    (salida / 'verificacion_local.json').write_text(json.dumps(resultado, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(resultado, ensure_ascii=False, indent=2), flush=True)


if __name__ == '__main__':
    ejecutar(sys.argv[1] if len(sys.argv) > 1 else RAIZ / 'entregas_locales' / 'actividad_20260911')
